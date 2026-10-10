import copy
import csv
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE = Path(__file__).parent / "template.xlsx"
GENERATED = Path(__file__).parent / "generated"

# Template rows
FIRST_DATA_ROW = 8   # linha modelo ímpar (fundo branco)
EVEN_MODEL_ROW = 9   # linha modelo par (fundo azul claro)
TOTAL_MODEL_ROW = 10 # linha de TOTAL
FOOTER_ROW = 12      # nota sobre DNP

# Template columns
COL_NUMBER = 1
COL_REFERENCE = 2
COL_QTY = 3
COL_DESCRIPTION = 4
COL_MPN = 5
COL_LCSC = 6
COL_PACKAGE = 7
COL_UNIT_PRICE = 8
COL_TOTAL = 9
LAST_COL = 9

# Others
REFERENCE_CHARS_PER_LINE = 19
DESCRIPTION_CHARS_PER_LINE = 60
EXTRA_LINE_HEIGHT = 14
DNP_COLOR = "FFC00000"

def csv_kicad_read(csv_path):
    """Reads the KiCad CSV and returns only the components that go into the BOM"""
    parts = []

    with open(csv_path, encoding="utf-8") as file:
        reader = csv.DictReader(file)

        for row in reader:
            if row["Exclude from BOM"] != "":
                continue
            parts.append(row)

    return parts

def terminal_ask(terminal_question, default=""):
    """Prompts for input in the terminal. If the answer is empty, uses the default"""
    if default != "":
        terminal_question = terminal_question + " [" + default + "]"

    terminal_answer = input(terminal_question + ": ").strip()

    if terminal_answer == "":
        return default
    return terminal_answer

def spread_fill_placeholders(ws, values):
    """Replaces {project}, {designer}... with the values entered"""
    for row in ws.iter_rows(min_row=1, max_row=FIRST_DATA_ROW - 1):
        for cell in row:
            if not isinstance(cell.value, str):
                continue

            for name, value in values.items():
                placeholder = "{" + name + "}"
                cell.value = cell.value.replace(placeholder, value)

def spread_copy_row_style(ws, row):
    """Stores the style of each cell in a template row"""
    styles = []

    for col in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col)
        styles.append(copy.copy(cell._style))

    height = ws.row_dimensions[row].height
    return styles, height

def spread_apply_row_style(ws, row, style):
    """Applies the style stored by spread_copy_row_style to a row"""
    styles, height = style

    for col in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col)
        cell._style = copy.copy(styles[col - 1])

    ws.row_dimensions[row].height = height

def spread_count_lines(text, chars_per_line):
    """Estimates how many lines a text will wrap to inside the cell"""
    lines = 1
    line_length = 0

    for word in text.split(" "):
        if line_length == 0:
            line_length = len(word)
        elif line_length + 1 + len(word) <= chars_per_line:
            line_length = line_length + 1 + len(word)
        else:
            lines = lines + 1
            line_length = len(word)

    return lines

def spread_row_height(reference, description, base_height):
    """Increases the row height when Reference or Description wrap onto multiple lines"""
    reference_lines = spread_count_lines(reference, REFERENCE_CHARS_PER_LINE)
    description_lines = spread_count_lines(description, DESCRIPTION_CHARS_PER_LINE)

    lines = max(reference_lines, description_lines)
    extra_lines = lines - 1
    return base_height + extra_lines * EXTRA_LINE_HEIGHT

def spread_paint_row_red(ws, row):
    """Colors a row’s text red (used for DNP components)"""
    for col in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col)
        font = copy.copy(cell.font)
        font.color = DNP_COLOR
        cell.font = font


def main():
    if len(sys.argv) != 2:
        print("Usage: python bom.py path/to/project.csv")
        sys.exit(1)

    csv_path = Path(sys.argv[1])
    xlsx_path = GENERATED / (csv_path.stem + "_BOM.xlsx")

    parts = csv_kicad_read(csv_path)

    total_parts = 0
    for part in parts:
        total_parts = total_parts + int(part["Qty"])

    print("Header fields (Enter keeps the value in brackets):")
    values = {
        "project"   : terminal_ask("Project", csv_path.stem),
        "designer"  : terminal_ask("Designed by"),
        "revision"  : terminal_ask("Revision"),
        "date"      : terminal_ask("Report date", date.today().strftime("%d-%b-%Y").upper()),
        "total"     : str(total_parts)
    }

    wb = load_workbook(TEMPLATE)
    ws = wb.active

    spread_fill_placeholders(ws, values)

    # Stores the styles of the template rows before deleting them
    odd_style = spread_copy_row_style(ws, FIRST_DATA_ROW)
    even_style = spread_copy_row_style(ws, EVEN_MODEL_ROW)
    total_style = spread_copy_row_style(ws, TOTAL_MODEL_ROW)

    footer = ws.cell(row=FOOTER_ROW, column=1)
    footer_text = footer.value
    footer_style = copy.copy(footer._style)

    ws.delete_rows(FIRST_DATA_ROW, ws.max_row)

    # One row for component
    row = FIRST_DATA_ROW
    for number, part in enumerate(parts, start=1):
        if number % 2 == 1:
            spread_apply_row_style(ws, row, odd_style)
        else:
            spread_apply_row_style(ws, row, even_style)

        # “C1,C2” becomes “C1, C2” so the text wraps between references
        reference = part["Reference"].replace(",", ", ")
        description = part["Description"]

        base_height = ws.row_dimensions[row].height
        ws.row_dimensions[row].height = spread_row_height(reference, description, base_height)

        ws.cell(row=row, column=COL_NUMBER, value=number)
        ws.cell(row=row, column=COL_REFERENCE, value=reference)
        ws.cell(row=row, column=COL_QTY, value=int(part["Qty"]))
        ws.cell(row=row, column=COL_DESCRIPTION, value=description)
        ws.cell(row=row, column=COL_MPN, value=part["MPN"])
        ws.cell(row=row, column=COL_LCSC, value=part["LCSC"])
        ws.cell(row=row, column=COL_PACKAGE, value=part["Package"])

        # Unit Price is left empty to be filled in by hand; Total = Qty × Unit Price
        total_formula = f'=IF(H{row}="","",C{row}*H{row})'
        ws.cell(row=row, column=COL_TOTAL, value=total_formula)

        if part["DNP"] != "":
            spread_paint_row_red(ws, row)

        row = row + 1

    last_data_row = row - 1

    spread_apply_row_style(ws, row, total_style)
    ws.cell(row=row, column=COL_REFERENCE, value="TOTAL")
    ws.cell(row=row, column=COL_QTY, value=f"=SUM(C{FIRST_DATA_ROW}:C{last_data_row})")
    ws.cell(row=row, column=COL_TOTAL, value=f"=SUM(I{FIRST_DATA_ROW}:I{last_data_row})")

    # Nota de rodapé, uma linha em branco abaixo do total
    footer_row = row + 2
    footer = ws.cell(row=footer_row, column=1, value=footer_text)
    footer._style = footer_style

    wb.save(xlsx_path)
    print("BOM saved to", xlsx_path)


if __name__ == "__main__":
    main()